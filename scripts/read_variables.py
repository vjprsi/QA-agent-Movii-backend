#!/usr/bin/env python3
"""Parse /docs/variables*.xlsx (+ Postman collection) into run-folder JSON files.

Outputs (in --out):
  variables.json          real values (local only)
  variables.masked.md     table for the report
  test_data.json          rows of the optional `test_data` sheet
  variables-check.json    missing_required, unresolved_postman_vars, deployed
Exit code is always 0 — missing items are reported, never fatal.
"""
import argparse
import glob
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _mask import mask, is_secret_key, is_pii_key  # noqa: E402

REQUIRED_ALWAYS = ["service_name", "deployed", "base_url", "k8s_namespace",
                   "db_host", "db_schema", "auth_type", "postman_collection"]
REQUIRED_IF = {
    "deploy_pipeline": lambda v: str(v.get("deployed", "")).strip().lower() in ("no", "false", "0", ""),
    "auth_token_url": lambda v: str(v.get("auth_type", "")).lower().startswith("oauth2"),
    "client_id": lambda v: str(v.get("auth_type", "")).lower().startswith("oauth2"),
    "client_secret": lambda v: str(v.get("auth_type", "")).lower().startswith("oauth2"),
}
STRESS_REQUIRED = ["stress_tps_ladder"]

# Postman variable name → candidate keys in variables.xlsx
POSTMAN_ALIASES = {
    "baseurl": "base_url", "base_url": "base_url", "host": "base_url", "url": "base_url",
    "token": "admin_token", "access_token": "admin_token", "bearer": "admin_token",
    "msisdn": "test_msisdn", "phone": "test_msisdn", "celular": "test_msisdn",
    "password": "test_password", "pin": "test_pin", "otp": "test_otp",
    "clientid": "client_id", "client_id": "client_id", "clientsecret": "client_secret", "client_secret": "client_secret",
}


def load_xlsx(path):
    try:
        import openpyxl  # noqa
    except ImportError:
        print("openpyxl missing → pip install openpyxl", file=sys.stderr)
        return {}, [], ["python package openpyxl not installed"]
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    variables, test_data, notes = {}, [], []

    def norm(s):
        return re.sub(r"[^a-z0-9_]", "", str(s or "").strip().lower().replace(" ", "_"))

    # variables sheet (first sheet if not named)
    vs = wb["variables"] if "variables" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(vs.iter_rows(values_only=True))
    if rows:
        header = [norm(h) for h in rows[0]]
        idx = {h: i for i, h in enumerate(header)}
        ki, vi = idx.get("key", 0), idx.get("value", 1)
        for r in rows[1:]:
            if not r or r[ki] in (None, ""):
                continue
            key = norm(r[ki])
            val = r[vi] if vi < len(r) else None
            variables[key] = "" if val is None else str(val).strip()
            meta = {}
            for extra in ("required", "category", "description"):
                if extra in idx and idx[extra] < len(r) and r[idx[extra]] not in (None, ""):
                    meta[extra] = str(r[idx[extra]]).strip()
            if meta:
                variables.setdefault("_meta", {})[key] = meta
    else:
        notes.append("variables sheet is empty")

    if "test_data" in wb.sheetnames:
        ts = wb["test_data"]
        trows = list(ts.iter_rows(values_only=True))
        if trows:
            th = [norm(h) for h in trows[0]]
            for r in trows[1:]:
                if not r or all(c in (None, "") for c in r):
                    continue
                test_data.append({th[i]: ("" if c is None else str(c).strip()) for i, c in enumerate(r) if i < len(th) and th[i]})
    return variables, test_data, notes


def postman_vars(collection_path):
    """Return set of {{var}} names used in the collection (minus those it defines itself)."""
    try:
        with open(collection_path, encoding="utf-8") as f:
            raw = f.read()
    except Exception as e:
        return set(), f"postman collection not readable: {e}"
    used = set(m.lower() for m in re.findall(r"{{\s*([A-Za-z0-9_.\-]+)\s*}}", raw))
    defined = set()
    try:
        col = json.loads(raw)
        for v in col.get("variable", []) or []:
            defined.add(str(v.get("key", "")).lower())
    except Exception:
        pass
    # dynamic postman vars are fine
    used = {u for u in used if not u.startswith("$")}
    return used - defined, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="path or glob to variables*.xlsx")
    ap.add_argument("--postman", help="Postman collection json")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    check = {"variables_file": None, "missing_required": [], "missing_for_stress": [],
             "unresolved_postman_vars": [], "deployed": "missing", "notes": []}

    candidates = sorted(glob.glob(a.xlsx)) if any(ch in a.xlsx for ch in "*?[") else [a.xlsx]
    candidates = [c for c in candidates if os.path.exists(c)]
    variables, test_data = {}, []
    if not candidates:
        check["notes"].append(f"variables file not found: {a.xlsx}")
    else:
        check["variables_file"] = candidates[0]
        variables, test_data, notes = load_xlsx(candidates[0])
        check["notes"] += notes

    meta = variables.pop("_meta", {})
    # defaults
    variables.setdefault("health_path", "/actuator/health")
    variables.setdefault("db_port", "3306")
    variables.setdefault("stress_step_seconds", "60")
    variables.setdefault("stress_error_rate_max", "0.01")
    if variables.get("stress_tps_ladder") and not variables.get("stress_max_tps"):
        try:
            variables["stress_max_tps"] = str(max(int(x) for x in re.split(r"[,\s;]+", variables["stress_tps_ladder"]) if x))
        except ValueError:
            check["notes"].append("stress_tps_ladder is not a list of integers")
    if variables.get("service_name") and not variables.get("k8s_deployment"):
        variables["k8s_deployment"] = variables["service_name"]

    dep = str(variables.get("deployed", "")).strip().lower()
    check["deployed"] = "yes" if dep in ("yes", "y", "true", "1", "si", "sí") else ("no" if dep in ("no", "n", "false", "0") else "missing")

    required = list(REQUIRED_ALWAYS) + [k for k, cond in REQUIRED_IF.items() if cond(variables)]
    # keys marked required=yes in the sheet are also required
    required += [k for k, m in meta.items() if str(m.get("required", "")).lower() in ("yes", "y", "true", "1")]
    check["missing_required"] = sorted({k for k in required if not variables.get(k)})
    check["missing_for_stress"] = [k for k in STRESS_REQUIRED if not variables.get(k)]

    if a.postman:
        used, err = postman_vars(a.postman)
        if err:
            check["notes"].append(err)
        keys_lower = {k.lower() for k in variables} | {f"test_data.{r.get('label','')}".lower() for r in test_data}
        unresolved = []
        for u in sorted(used):
            cand = POSTMAN_ALIASES.get(u, u)
            if u in keys_lower or cand in keys_lower or u.replace("-", "_") in keys_lower:
                continue
            unresolved.append(u)
        check["unresolved_postman_vars"] = unresolved

    with open(os.path.join(a.out, "variables.json"), "w", encoding="utf-8") as f:
        json.dump(variables, f, indent=2, ensure_ascii=False)
    with open(os.path.join(a.out, "test_data.json"), "w", encoding="utf-8") as f:
        json.dump(test_data, f, indent=2, ensure_ascii=False)
    with open(os.path.join(a.out, "variables-check.json"), "w", encoding="utf-8") as f:
        json.dump(check, f, indent=2, ensure_ascii=False)

    lines = ["| key | value | category |", "|---|---|---|"]
    for k in sorted(variables):
        v = variables[k]
        shown = mask(v) if (is_secret_key(k) or is_pii_key(k)) else v
        lines.append(f"| {k} | `{shown}` | {meta.get(k, {}).get('category', '')} |")
    if test_data:
        lines += ["", f"test_data rows: {len(test_data)} (labels: " + ", ".join(r.get("label", "?") for r in test_data) + ")"]
    with open(os.path.join(a.out, "variables.masked.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    # console summary (masked)
    print(json.dumps({k: v for k, v in check.items()}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
