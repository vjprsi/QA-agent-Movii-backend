#!/usr/bin/env python3
"""Produce Word (.docx) copies of a run's JSON results so the QA team can read
and assemble a final document without opening JSON.

Writes into  docs/<service>/<run>/word/ :
  - results-summary.docx : one table per results/*.json suite (id, title,
    status, severity, observed) + a counts table.
  - evidence.docx        : for every evidence/<ID>/reqN.json + respN.json, the
    request (method, URL, headers, body) and response (status, latency,
    headers, body) as readable blocks.
  - stress-results.docx  : stress/summary.json rendered as tables, plus the
    text of stress/stress-results.md if present.

All inputs are already masked (qa_client masks secrets/PII before writing
evidence); this script only reformats them, it does not unmask anything.

Usage:
  python -I scripts/results_to_docx.py --run-dir docs/<service>/<run>
  (optional --service NAME ; defaults derived from the path)
"""
import argparse
import glob
import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

MONO = "Consolas"
SEV_COLOR = {"CRITICAL": "B71C1C", "HIGH": "E65100", "MEDIUM": "F9A825", "LOW": "2E7D32"}
STATUS_COLOR = {"PASS": "1B5E20", "FAIL": "B71C1C", "BLOCKED": "6A1B9A",
                "NOT RUN": "616161", "UNVERIFIED": "6A1B9A"}


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def mono_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(6)
    run = p.add_run(text)
    run.font.name = MONO
    run.font.size = Pt(9)
    return p


def pretty(obj):
    try:
        return json.dumps(obj, indent=2, ensure_ascii=False)
    except Exception:
        return str(obj)


def colored_cell(cell, text, hexcolor=None, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    if hexcolor:
        run.font.color.rgb = RGBColor.from_string(hexcolor)


def add_table(doc, headers, rows, color_map=None, color_col=None, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        colored_cell(t.rows[0].cells[j], h, bold=True)
    for r in rows:
        cells = t.add_row().cells
        for j, val in enumerate(r):
            if color_map and color_col == j:
                colored_cell(cells[j], val, color_map.get(str(val).upper()), bold=True)
            else:
                cells[j].text = "" if val is None else str(val)
                for par in cells[j].paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(9)
    return t


def meta_line(doc, service, run):
    p = doc.add_paragraph()
    r = p.add_run(f"{service}   |   run {run}")
    r.italic = True
    r.font.color.rgb = RGBColor.from_string("555555")


def results_summary(run_dir, service, run, out_dir):
    results = sorted(glob.glob(os.path.join(run_dir, "results", "*.json")))
    if not results:
        return None
    doc = Document()
    doc.add_heading(f"Test results — {service}", 0)
    meta_line(doc, service, run)
    grand = {}
    for rp in results:
        data = json.load(open(rp, encoding="utf-8"))
        if not isinstance(data, list):
            continue
        heading(doc, os.path.basename(rp), 1)
        counts = {}
        rows = []
        for c in data:
            st = str(c.get("status", "")).upper()
            counts[st] = counts.get(st, 0) + 1
            grand[st] = grand.get(st, 0) + 1
            rows.append([c.get("id", ""), c.get("title", ""), st,
                         c.get("severity") or "", c.get("observed", "")])
        add_table(doc, ["Counts"] + list(counts.keys()), [["cases"] + list(counts.values())])
        doc.add_paragraph()
        add_table(doc, ["ID", "Title", "Status", "Severity", "Observed"], rows,
                  color_map=STATUS_COLOR, color_col=2)
        doc.add_paragraph()
    heading(doc, "Overall", 1)
    add_table(doc, list(grand.keys()) or ["(none)"],
              [list(grand.values()) or [0]])
    out = os.path.join(out_dir, "results-summary.docx")
    doc.save(out)
    return out


def evidence_doc(run_dir, service, run, out_dir):
    evid = os.path.join(run_dir, "evidence")
    if not os.path.isdir(evid):
        return None
    doc = Document()
    doc.add_heading(f"Request / response evidence — {service}", 0)
    meta_line(doc, service, run)
    ids = sorted(os.listdir(evid))
    for tid in ids:
        d = os.path.join(evid, tid)
        if not os.path.isdir(d):
            continue
        reqs = sorted(glob.glob(os.path.join(d, "req*.json")),
                      key=lambda p: int(re.search(r"(\d+)", os.path.basename(p)).group(1)))
        if not reqs:
            continue
        heading(doc, tid, 1)
        for rq in reqs:
            n = re.search(r"(\d+)", os.path.basename(rq)).group(1)
            rsp = os.path.join(d, f"resp{n}.json")
            try:
                req = json.load(open(rq, encoding="utf-8"))
            except Exception:
                continue
            heading(doc, f"Request {n}", 3)
            line = f"{req.get('method','')} {req.get('url','')}"
            mono_block(doc, line)
            if req.get("headers"):
                mono_block(doc, "headers: " + pretty(req["headers"]))
            if req.get("body") is not None:
                mono_block(doc, "body: " + pretty(req["body"]))
            if os.path.exists(rsp):
                resp = json.load(open(rsp, encoding="utf-8"))
                heading(doc, f"Response {n}", 3)
                status = resp.get("status", "")
                p = doc.add_paragraph()
                r = p.add_run(f"HTTP {status}   {resp.get('ms','')} ms")
                r.bold = True
                if str(status).startswith(("4", "5")):
                    r.font.color.rgb = RGBColor.from_string("B71C1C")
                else:
                    r.font.color.rgb = RGBColor.from_string("1B5E20")
                if resp.get("headers"):
                    mono_block(doc, "headers: " + pretty(resp["headers"]))
                if resp.get("body") is not None:
                    mono_block(doc, "body: " + pretty(resp["body"]))
        doc.add_paragraph()
    out = os.path.join(out_dir, "evidence.docx")
    doc.save(out)
    return out


def stress_doc(run_dir, service, run, out_dir):
    sdir = os.path.join(run_dir, "stress")
    summ = os.path.join(sdir, "summary.json")
    if not os.path.exists(summ):
        return None
    data = json.load(open(summ, encoding="utf-8"))
    doc = Document()
    doc.add_heading(f"Stress test — {service}", 0)
    meta_line(doc, service, run)
    top = [[k, v] for k, v in data.items() if not isinstance(v, dict)]
    heading(doc, "Run parameters", 1)
    add_table(doc, ["Metric", "Value"], top)
    metric_keys = ["n", "p50", "p95", "p99", "max", "mean", "status"]
    groups = [(k, v) for k, v in data.items() if isinstance(v, dict)]
    if groups:
        heading(doc, "Latency by endpoint (ms)", 1)
        rows = []
        for name, g in groups:
            rows.append([name] + [g.get(mk, "") for mk in metric_keys])
        add_table(doc, ["Group"] + metric_keys, rows)
    md = os.path.join(sdir, "stress-results.md")
    if os.path.exists(md):
        heading(doc, "Notes", 1)
        for line in open(md, encoding="utf-8").read().splitlines():
            if line.strip():
                doc.add_paragraph(re.sub(r"[#*`]", "", line))
    out = os.path.join(out_dir, "stress-results.docx")
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run-dir", required=True, help="docs/<service>/<run>")
    ap.add_argument("--service")
    args = ap.parse_args()
    run_dir = os.path.normpath(args.run_dir)
    parts = run_dir.replace("\\", "/").split("/")
    service = args.service or (parts[-2] if len(parts) >= 2 else "service")
    run = parts[-1]
    out_dir = os.path.join(run_dir, "word")
    os.makedirs(out_dir, exist_ok=True)
    written = []
    for fn in (results_summary, evidence_doc, stress_doc):
        try:
            p = fn(run_dir, service, run, out_dir)
            if p:
                written.append(p)
        except Exception as e:  # noqa: BLE001
            print(f"skip {fn.__name__}: {e}")
    for p in written:
        print("wrote", p)
    if not written:
        print("no JSON inputs found under", run_dir)


if __name__ == "__main__":
    main()
